#!/usr/bin/env python3
"""Rebuild a source-backed, editable Thin Signals deck in the official template."""
from __future__ import annotations

import argparse
import json
import zipfile
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / 'Thin-Signals-2026-09-29.pptx'
TEMPLATE_SHA256 = '1604aa9146b82cacd96f893be2c5e08c6981183ba155cd4a11c245f25cd894d5'
EVENT_DATE = datetime(2026, 9, 29, 12, 0, tzinfo=UTC)

WHITE = 'FFFFFF'
INK = '1C1D22'
DEEP = '310F53'
PURPLE = '520978'
PINK = 'FF0053'
LAV = '8A83D1'
MUTED = '64616C'
LIGHT = 'F4F0F7'
BORDER = 'E5DEEB'
PALE = 'FFD6E4'
TEAL = '237A71'
FONT = 'Montserrat'


def color(value):
    return RGBColor.from_string(value)


def set_text(shape, content, size=16, fill=INK, bold=False, align=PP_ALIGN.LEFT,
             valign=MSO_ANCHOR.TOP, margin=0.0, spacing=1.13, after=5):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = valign
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    for i, value in enumerate(content.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_before = Pt(0)
        p.space_after = Pt(after)
        # Suppress inherited bullets/margins without changing the official box.
        ppr = p._p.get_or_add_pPr()
        ppr.set('marL', '0')
        ppr.set('indent', '0')
        for item in list(ppr):
            if item.tag.rsplit('}', 1)[-1] in {'buChar', 'buAutoNum', 'buBlip', 'buNone'}:
                ppr.remove(item)
        ppr.append(OxmlElement('a:buNone'))
        r = p.add_run()
        r.text = value
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color(fill)
    return shape


def text(slide, x, y, w, h, content, **kw):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    return set_text(sh, content, **kw)


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color(fill)
    if line:
        sh.line.color.rgb = color(line)
        sh.line.width = Pt(0.8)
    else:
        sh.line.fill.background()
    if radius:
        sh.adjustments[0] = 0.12
    return sh


def pill(slide, x, y, w, content, fill=LIGHT, ink=PURPLE):
    rect(slide, x, y, w, 0.31, fill=fill, radius=True)
    text(slide, x + 0.08, y + 0.035, w - 0.16, 0.25, content,
         size=9, fill=ink, bold=True, align=PP_ALIGN.CENTER, spacing=1, after=0)


def card(slide, x, y, w, h, number, title, body, accent=PINK, dark=False):
    rect(slide, x, y, w, h, fill=DEEP if dark else WHITE,
         line=None if dark else BORDER, radius=True)
    rect(slide, x, y + 0.23, 0.045, 0.58, fill=accent)
    text(slide, x + 0.22, y + 0.23, w - 0.44, 0.48, number,
         size=28, fill=accent, bold=True, after=0)
    text(slide, x + 0.22, y + 0.88, w - 0.44, 0.65, title,
         size=19, fill=WHITE if dark else DEEP, bold=True)
    text(slide, x + 0.22, y + 1.57, w - 0.44, h - 1.67, body,
         size=14, fill='E9DFF1' if dark else MUTED)


def banner(slide, message, *, y=6.26, fill=DEEP, ink=WHITE, height=0.56, size=14):
    rect(slide, 0.58, y, 12.17, height, fill=fill, radius=True)
    text(slide, 0.78, y + 0.13, 11.77, height - 0.19, message,
         size=size, fill=ink, bold=True, after=0, spacing=1.06)


def note(slide, content):
    slide.notes_slide.notes_text_frame.text = content


def new_slide(prs, number, eyebrow, title, subtitle='', dark=False):
    blank = next(x for x in prs.slide_layouts if x.name == 'Пустой слайд')
    slide = prs.slides.add_slide(blank)
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)
    slide._element.set('showMasterSp', '0')
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(DEEP if dark else WHITE)
    rect(slide, 0, 0, 13.33333, 0.065, PINK)
    text(slide, 0.58, 0.36, 11.6, 0.25, eyebrow.upper(),
         size=10, fill=PALE if dark else PURPLE, bold=True, after=0)
    text(slide, 0.58, 0.89, 12.05, 0.99, title,
         size=29, fill=WHITE if dark else DEEP, bold=True, spacing=1.03, after=0)
    if subtitle:
        text(slide, 0.60, 1.96, 11.98, 0.58, subtitle,
             size=15, fill='E9DFF1' if dark else MUTED, spacing=1.12, after=0)
    rect(slide, 0.58, 7.02, 12.17, 0.008, fill=LAV if dark else BORDER)
    text(slide, 0.58, 7.16, 8.4, 0.16,
         'AGI SPACE  /  ТОНКИЕ СИГНАЛЫ  ·  КОНЦЕПТ 29.09.2026',
         size=7.8, fill='E9DFF1' if dark else MUTED, bold=True, after=0)
    text(slide, 11.55, 7.12, 1.2, 0.22, f'{number:02d} / 15',
         size=10, fill=PALE if dark else PURPLE, bold=True,
         align=PP_ALIGN.RIGHT, after=0)
    return slide


def native_picture(ph, path, max_scale=0.92):
    left, top, width, height = ph.left, ph.top, ph.width, ph.height
    with Image.open(path) as im:
        ratio = im.width / im.height
    w = min(width * max_scale, height * max_scale * ratio)
    h = w / ratio
    pic = ph.insert_picture(str(path))
    pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0
    pic.left = int(left + (width - w) / 2)
    pic.top = int(top + (height - h) / 2)
    pic.width, pic.height = int(w), int(h)
    return pic


def source_note(*sources):
    return '\n'.join([
        'Концепт / исследовательские материалы: 29.09.2026.',
        'Дата комплекта 29.09.2026 установлена по разрешению организаторов, сообщённому оператором.',
        'Дата комплекта не подтверждает конкурсную сдачу или дату релиза.',
        'Источники:', *sources,
    ])


def mandatory_base(base):
    """Preserve the five accepted organizer slides as editable source objects."""
    prs = Presentation(base)
    assert len(prs.slides) >= 5
    for i in reversed(range(5, len(prs.slides))):
        sid = prs.slides._sldIdLst[i]
        prs.part.drop_rel(sid.rId)
        del prs.slides._sldIdLst[i]
    for slide in prs.slides:
        old_note = slide.notes_slide.notes_text_frame.text
        old_note = '\n'.join(
            line for line in old_note.split('\n')
            if not line.startswith(('Фактическая генерация и актуализация реализации:',
                                    'Актуализация реализации и проверка презентации:'))
        )
        note(slide, old_note.replace('Эти даты не подтверждают', 'Дата комплекта не подтверждает'))
    # Roundtrip resolves slide-part numbering after deleting the old content.
    stream = BytesIO()
    prs.save(stream)
    stream.seek(0)
    return Presentation(stream)


def solution_slides(prs, screenshot: Path | None):
    s = new_slide(prs, 6, 'Проблема', 'Раннее изменение теряется\nмежду новостью и решением',
                  'Задача аналитика — найти новое техническое основание и проверить, почему оно важно сейчас.')
    card(s, 0.58, 2.87, 3.87, 2.96, '01', 'Информационный шум',
         'Одна публикация расходится по агрегаторам. Число упоминаний не равно числу независимых фактов.')
    card(s, 4.73, 2.87, 3.87, 2.96, '02', 'Ложная новизна',
         'Известный рынок или зрелая категория могут выглядеть как новый сигнал без проверки контрпризнаков.', accent=PURPLE)
    card(s, 8.88, 2.87, 3.87, 2.96, '03', 'Потерянный задел',
         'Старое направление становится интересным, когда меняются материалы, производство, стоимость или спрос.', accent=LAV)
    banner(s, 'Ценность продукта: меньше времени на первичный разбор, больше внимания решающей проверке.')
    note(s, source_note('CONCEPT.md', 'docs/METHODOLOGY.md', 'Исследовательская гипотеза ценности: экономия времени пока не измерена.'))

    s = new_slide(prs, 7, 'Продукт', 'Три модуля — три разных решения',
                  'Общий принцип: сохранить объяснение, источник и границу уверенности.')
    card(s, 0.58, 2.85, 3.87, 3.12, '01', 'Тонкие сигналы',
         'Живой поиск по запросу.\n\nПризнаки, оценка, допуск, наблюдение и причины исключения.')
    card(s, 4.73, 2.85, 3.87, 3.12, '02', 'Потенциал',
         'Кандидаты ниже порога.\n\nБлизость к допуску или выраженный потенциал — повод для следующей проверки.', accent=PURPLE)
    card(s, 8.88, 2.85, 3.87, 3.12, '03', 'Новая жизнь',
         'Поиск по архиву задела.\n\nСтарый барьер → новые условия → гипотеза продукта → сравнительный тест.', accent=LAV)
    banner(s, 'Три очереди сохраняют собственный смысл: наблюдение, приоритет проверки и исследовательский архив.')
    note(s, source_note('docs/ui/EVIDENCE_FIRST_V3.md', 'ts/desk.py', 'ts/potential.py', 'ts/old_tech.py'))

    s = new_slide(prs, 8, '01 / Тонкие сигналы', 'От вопроса — к объяснённому результату',
                  'До 15 карточек в выдаче. Число результатов определяется найденными основаниями.')
    labels = [
        ('1', 'Запрос', 'Направление и контекст\nобычными словами'),
        ('2', 'Источники', 'Поиск, дедупликация,\nчтение фрагментов'),
        ('3', 'Проверки', 'P01–P18, M01–M06,\nпереход стадии'),
        ('4', 'Результат', 'Вердикт, оценка,\nисточники и причины'),
    ]
    for i, (num, heading, body) in enumerate(labels):
        x = 0.58 + i * 3.10
        rect(s, x, 2.92, 2.86, 1.98, fill=LIGHT, radius=True)
        pill(s, x + 0.18, 3.10, 0.40, num, fill=PALE, ink=PINK)
        text(s, x + 0.18, 3.58, 2.48, 0.35, heading, size=18, fill=DEEP, bold=True)
        text(s, x + 0.18, 4.13, 2.48, 0.57, body, size=12.6, fill=MUTED)
        if i < 3:
            text(s, x + 2.87, 3.64, 0.23, 0.35, '›', size=24, fill=PINK, bold=True)
    rect(s, 0.58, 5.27, 12.17, 0.63, fill=DEEP, radius=True)
    text(s, 0.80, 5.46, 11.73, 0.28,
         'rules-0.1:  <0,35 — не допущено  /  0,35 ≤ балл <0,60 — наблюдение  /  ≥0,60 — слабый сигнал',
         size=12.4, fill=WHITE, bold=True, after=0)
    banner(s, 'Позиция в списке не меняет вердикт: «наблюдение» остаётся наблюдением.', fill=LIGHT, ink=DEEP)
    note(s, source_note('ts/scoring/rules-0.1.yaml: thresholds (ledger/live)', 'ts/desk.py: вердикты',
                        'Пороги description-mode отличаются; на этом слайде показан основной live/ledger-контракт.',
                        'Поля даты пока задают контекст запроса, а не гарантированный типизированный API-фильтр.'))

    s = new_slide(prs, 9, '02 / Потенциал', 'Не прошёл порог —\nможет заслуживать следующей проверки',
                  'Отдельная очередь повторного анализа; основной допуск и его правила остаются неизменными.')
    rect(s, 0.58, 2.86, 5.04, 2.97, fill=DEEP, radius=True)
    text(s, 0.87, 3.13, 4.46, 0.27, 'КАК КАНДИДАТ ПОПАДАЕТ В ОЧЕРЕДЬ', size=10, fill=PALE, bold=True)
    text(s, 0.87, 3.70, 4.46, 0.56, 'Близок к порогу', size=23, fill=WHITE, bold=True)
    pill(s, 0.88, 4.35, 0.65, 'ИЛИ', fill=PINK, ink=WHITE)
    text(s, 0.87, 4.96, 4.46, 0.58, 'Высокий потенциал\nпо найденным признакам', size=20, fill=WHITE, bold=True, spacing=1.07)
    for y, label, sub, shade in [
        (2.91, 'Технический', 'Новый механизм, эффективность, применение', PINK),
        (3.92, 'Спрос', 'Пользователь, стимул, ценность', PURPLE),
        (4.93, 'Развитие', 'Экосистема, конвергенция, исследовательский задел', LAV),
    ]:
        rect(s, 5.98, y, 6.77, 0.87, fill=LIGHT, radius=True)
        rect(s, 5.98, y + 0.15, 0.045, 0.53, fill=shade)
        text(s, 6.22, y + 0.12, 6.28, 0.26, label, size=16, fill=DEEP, bold=True)
        text(s, 6.22, y + 0.50, 6.28, 0.24, sub, size=11.5, fill=MUTED, after=0)
    banner(s, 'Шкала 0–100 означает приоритет проверки, а не вероятность спроса, внедрения или доходности.')
    note(s, source_note('ts/potential.py', 'ts/desk.py: potential_board',
                        'potential-v2: near-watchlist 0.25–0.35; high-potential из watchlist 0.35–0.60. Только текущие scored-карточки, без запрещающих контрпризнаков.',
                        'Неизвестные измерения остаются None; ноль не подставляется вместо отсутствия данных.',
                        'Формула исследовательская и детерминированная, не валидированная рыночная прогностическая модель.'))

    s = new_slide(prs, 10, '03 / Новая жизнь', 'Старый задел. Новые условия.\nПроверяемое применение.',
                  'Поиск по зафиксированному исследовательскому архиву от 29.09.2026 — не свежий веб-поиск.')
    rect(s, 0.58, 2.83, 12.17, 0.67, fill=DEEP, radius=True)
    text(s, 0.83, 3.00, 11.67, 0.34, 'Пример из архива: насосы Кнудсена для газоанализа на микрочипе',
         size=18, fill=WHITE, bold=True)
    cols = [
        ('Старый барьер', 'Сложно изготовить узкие каналы; малая производительность и теплопотери.'),
        ('Что изменилось', 'MEMS-процессы и интеграция компонентов в газоаналитический чип.'),
        ('Гипотеза продукта', 'Модуль переносного газоанализатора. Требует сравнения с микронасосом.'),
    ]
    for i, (heading, body) in enumerate(cols):
        x = 0.58 + i * 4.15
        rect(s, x, 3.77, 3.87, 1.65, fill=LIGHT, radius=True)
        text(s, x + 0.20, 3.96, 3.47, 0.32, heading, size=16, fill=DEEP, bold=True)
        text(s, x + 0.20, 4.48, 3.47, 0.74, body, size=12.8, fill=MUTED)
    text(s, 0.63, 5.73, 2.83, 0.4, '15 кандидатов', size=21, fill=PURPLE, bold=True)
    text(s, 3.62, 5.79, 8.88, 0.37,
         '0 полностью прошли коммерческую и правовую квалификацию', size=16, fill=MUTED, bold=True)
    banner(s, 'Выход: первый продукт, возможный покупатель, риски, открытые вопросы и решающий эксперимент.',
           fill=LIGHT, ink=DEEP)
    note(s, source_note('data/old_tech/snapshot_2026-09-29.json, T01, S04/S05/S06/S82',
                        'Архивная карточка: https://pmc.ncbi.nlm.nih.gov/articles/PMC12686500/',
                        'Патентный след: https://patents.google.com/patent/US9243624B2/en',
                        'Техническое свидетельство, рыночная гипотеза и права проверяются отдельно.'))

    s = new_slide(prs, 11, 'Доверие к результату', 'У каждого вывода — свой уровень основания',
                  'Карточка разделяет то, что найдено в источниках, и то, что ещё предстоит проверить.')
    groups = [
        ('Факты', 'Доступные фрагменты\nи их ссылки', PINK),
        ('Гипотезы', 'Применение, выгода\nи открытые вопросы', PURPLE),
        ('Проверки', 'Признаки, контрпризнаки,\nвклад в оценку', LAV),
        ('Источники', 'Документ, дата, тип\nи происхождение', TEAL),
    ]
    for i, (heading, body, shade) in enumerate(groups):
        x = 0.58 + i * 3.10
        rect(s, x, 2.89, 2.86, 1.90, fill=LIGHT, radius=True)
        rect(s, x + 0.19, 3.10, 0.47, 0.045, fill=shade)
        text(s, x + 0.19, 3.35, 2.48, 0.34, heading, size=19, fill=DEEP, bold=True)
        text(s, x + 0.19, 3.94, 2.48, 0.51, body, size=12.8, fill=MUTED)
    text(s, 0.70, 5.02, 11.8, 0.42,
         'Фрагмент  →  признак  →  вклад в оценку  →  вердикт  →  решение аналитика',
         size=18, fill=DEEP, bold=True, align=PP_ALIGN.CENTER)
    text(s, 0.75, 5.58, 11.74, 0.32,
         'Модель структурирует материал. Код рассчитывает балл. Человек оценивает полезность.',
         size=13.5, fill=MUTED, align=PP_ALIGN.CENTER)
    banner(s, 'Статус verified требует ссылки на переданный фрагмент; неподтверждённый вывод не повышает оценку.')
    note(s, source_note('docs/METHODOLOGY.md', 'STATUS.md: версия проверяется по frozen candidate', 'docs/ui/EVIDENCE_FIRST_V3.md',
                        'Проверка валидной ссылки необходима, но сама по себе не является доказательством независимой репликации.',
                        'Полный долговечный evidence ledger остаётся открытой задачей.'))

    s = new_slide(prs, 12, 'Архитектура', 'Единый путь от источника до карточки',
                  'В рабочем контуре: FastAPI, Streamlit и PostgreSQL; версии правил и чтения сохранены.')
    blocks = [
        ('Запрос', 'API + план поиска\nтемы и контекст'),
        ('Сбор', 'Коннекторы\nдедупликация'),
        ('Чтение', 'Фрагменты\nLLM-чек-лист'),
        ('Решение', 'Правила\nоценка и причины'),
        ('Аналитик', 'История + UI\nтри модуля'),
    ]
    for i, (heading, body) in enumerate(blocks):
        x = 0.58 + i * 2.48
        rect(s, x, 2.94, 2.25, 1.89, fill=DEEP if i == 3 else LIGHT, radius=True)
        text(s, x + 0.18, 3.20, 1.89, 0.34, heading, size=17.2,
             fill=WHITE if i == 3 else DEEP, bold=True)
        text(s, x + 0.18, 3.88, 1.89, 0.57, body, size=12.2,
             fill='E9DFF1' if i == 3 else MUTED)
        if i < 4:
            text(s, x + 2.26, 3.63, 0.22, 0.40, '›', size=24, fill=PINK, bold=True)
    for i, (heading, body) in enumerate([
        ('Версии', 'Правила, чек-лист, модель и провайдер'),
        ('Воспроизводимость', 'Сохранённый прогон и причины решения'),
        ('Граница', 'Статус поиска не равен качеству результата'),
    ]):
        x = 0.64 + i * 4.13
        text(s, x, 5.22, 3.72, 0.28, heading, size=14, fill=DEEP, bold=True)
        text(s, x, 5.70, 3.72, 0.45, body, size=12.3, fill=MUTED)
    banner(s, 'Проверены ссылки на фрагменты, модель/провайдер и кэш; rules-0.1 и пороги сохранены.',
           fill=LIGHT, ink=DEEP)
    note(s, source_note('ARCHITECTURE.md', 'STATUS.md: open audit findings', 'docs/PROCESSING_CONTRACT.md',
                        'Долговечный per-step/evidence ledger реализован не полностью. Полный отказ коннекторов может завершиться succeeded — открытый дефект контракта.'))

    s = new_slide(prs, 13, 'Качество и готовность', 'Качество измеряется\nна независимых данных',
                  'Реализация правил и интерфейса не заменяет оценку полного поискового конвейера.')
    rect(s, 0.58, 2.85, 5.89, 2.80, fill=LIGHT, radius=True)
    rect(s, 6.78, 2.85, 5.97, 2.80, fill=DEEP, radius=True)
    text(s, 0.83, 3.13, 5.40, 0.29, 'РЕАЛИЗОВАНО И ПРОВЕРЕНО', size=11, fill=PURPLE, bold=True)
    text(s, 0.83, 3.70, 5.40, 1.55,
         'Три модуля и объяснение карточки\nДетерминированные правила\nПроверки ссылок, версий и кэша\nВоспроизводимый архив «Новая жизнь»',
         size=16, fill=DEEP, spacing=1.15, after=9)
    text(s, 7.03, 3.13, 5.47, 0.29, 'ЕЩЁ НЕ ПОДТВЕРЖДЕНО', size=11, fill=PALE, bold=True)
    text(s, 7.03, 3.70, 5.47, 1.55,
         'Accuracy ≥75% полного live-поиска\nНадёжный TOP-15 на новых запросах\nПолнота журнала доказательств\nФакт конкурсной сдачи',
         size=16, fill=WHITE, spacing=1.15, after=9)
    text(s, 0.65, 5.98, 12.02, 0.67,
         '100 примеров организаторов — seed/design set, а не нетронутый holdout.\nНужны независимая разметка, зафиксированная версия и end-to-end оценка.',
         size=15, fill=PURPLE, bold=True, spacing=1.15, after=0)
    note(s, source_note('AGENTS.md: данные и запрет заявлений без eval', 'STATUS.md', 'docs/EVAL.md (исторический description-mode)',
                        'Историческое F1 0.9541 не вынесено в результат: ранний эксперимент на известном наборе не доказывает качество текущего live-поиска.',
                        'Факт сдачи организаторам не подтверждён.'))

    s = new_slide(prs, 14, 'Интерфейс', 'Фильтр → рейтинг → проверяемое основание')
    with Image.open(screenshot) as im:
        ratio = im.width / im.height
    area = (0.58, 2.00, 9.40, 4.64)
    width = min(area[2], area[3] * ratio)
    height = width / ratio
    left = area[0] + (area[2] - width) / 2
    top = area[1] + (area[3] - height) / 2
    rect(s, left - 0.02, top - 0.02, width + 0.04, height + 0.04, fill=BORDER)
    pic = s.shapes.add_picture(str(screenshot), Inches(left), Inches(top),
                               width=Inches(width), height=Inches(height))
    pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0
    for y, heading, body, shade in [
        (2.10, 'Три модуля', 'Сигналы\nПотенциал\nНовая жизнь', PINK),
        (3.73, 'Основания', 'Факты и гипотезы\nПроверки и ссылки', PURPLE),
        (5.16, 'Контекст', 'Статус прогона\nи версия данных', LAV),
    ]:
        rect(s, 10.19, y, 0.045, 0.43, fill=shade)
        text(s, 10.40, y, 2.22, 0.34, heading, size=15, fill=DEEP, bold=True)
        text(s, 10.40, y + 0.48, 2.22, 0.83, body, size=12.5, fill=MUTED)
    text(s, 0.61, 6.77, 12.00, 0.19,
         'Текущая версия · комплект 29.09.2026 · снимок без обрезки и изменения пропорций',
         size=8.2, fill=MUTED, after=0)
    note(s, source_note('assets/current-three-panels.png; происхождение в release-evidence.json',
                        'Текущий код приложения, проверенный на изолированном стенде с контролируемой сессией.',
                        'Скриншот не является доказательством live-точности или авторизованной проверки production.'))

    s = new_slide(prs, 15, 'Следующий шаг', 'Раньше заметить.\nТочнее выбрать, что проверять.',
                  'Пилот с аналитиками должен измерить полезность карточек, а не только количество находок.', dark=True)
    cols = [
        ('01', 'Независимая оценка', 'Новые запросы, внешняя разметка, качество полного конвейера.'),
        ('02', 'Надёжность', 'Долговечный журнал, статусы отказа источников и монотонность оценки.'),
        ('03', 'Пилотный сценарий', 'Сократить свежий поиск и измерить время работы аналитика на реальных задачах.'),
    ]
    for i, (num, heading, body) in enumerate(cols):
        x = 0.58 + i * 4.15
        rect(s, x, 2.89, 3.87, 2.61, fill=WHITE, radius=True)
        text(s, x + 0.22, 3.12, 3.43, 0.42, num, size=28, fill=PINK, bold=True)
        text(s, x + 0.22, 3.83, 3.43, 0.36, heading, size=16, fill=DEEP, bold=True)
        text(s, x + 0.22, 4.49, 3.43, 0.73, body, size=13.2, fill=MUTED)
    text(s, 0.61, 5.92, 11.85, 0.66,
         'Рабочий контур: tsignal.aispaces.ru · три модуля доступны.\nСледующий результат — измеримая польза для аналитика.',
         size=18, fill=WHITE, bold=True, spacing=1.14, after=0)
    note(s, source_note('План развития, не обязательство заказчика и не подтверждённая поставка.',
                        'Данные команды на обязательных слайдах 2–4 требуют заполнения до внешней подачи.',
                        'Дата комплекта: 29.09.2026.',
                        'Свежий поиск: production ≈8:13, предыдущий сопоставимый staging ≈8:51. Это единичные реконструированные замеры, не SLA и не сравнение с прежним production.',
                        'Production: 48 источников, 2 watchlist и 2 potential; staging: 48 источников, 4 watchlist и 3 potential. Повтор по кэшу занимает доли секунды, не является новым поиском.',
                        'Авторизованный production smoke не заявляется; автоматизированный UI smoke использует изолированный стенд.',
                        'Монотонность rules-0.1, статусы полного отказа источников и полнота evidence ledger остаются открытыми задачами.'))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--base', type=Path, default=OUTPUT,
                        help='Editable accepted deck; slides 1–5 are retained exactly.')
    parser.add_argument('--screenshot', type=Path,
                        default=ROOT / 'assets/current-three-panels.png')
    parser.add_argument('--output', type=Path, default=OUTPUT)
    args = parser.parse_args()
    if not args.base.is_file() or not args.screenshot.is_file():
        parser.error('Both editable base deck and current screenshot must exist.')
    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs = mandatory_base(args.base)
    solution_slides(prs, args.screenshot)
    assert len(prs.slides) == 15
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.0001
    props = prs.core_properties
    props.title = 'Тонкие сигналы — ЛЦТ 2026'
    props.subject = 'Три модуля технологической разведки по проверяемым источникам'
    props.author = 'AGI Space'
    props.keywords = 'ЛЦТ 2026, Газпромбанк.Тех, AGI Space, дата комплекта 29.09.2026'
    props.comments = (
        'Дата комплекта 29.09.2026 разрешена организаторами со слов оператора. '
        'Не подтверждает конкурсную сдачу.')
    props.created = props.modified = EVENT_DATE
    props.last_modified_by = 'AGI Space'
    stream = BytesIO()
    prs.save(stream)
    stream.seek(0)
    # OOXML's ZIP container also carries dates, independently of core.xml.
    with zipfile.ZipFile(stream) as source, zipfile.ZipFile(args.output, 'w') as target:
        for entry in source.infolist():
            entry.date_time = EVENT_DATE.timetuple()[:6]
            entry.extra = b''
            entry.comment = b''
            target.writestr(entry, source.read(entry.filename))
    outline = []
    for i, slide in enumerate(prs.slides, 1):
        outline.append({'slide': i, 'layout': slide.slide_layout.name,
                        'texts': [sh.text for sh in slide.shapes if sh.has_text_frame and sh.text],
                        'notes': slide.notes_slide.notes_text_frame.text})
    (args.output.parent / '2026-09-29-outline.json').write_text(
        json.dumps(outline, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'output': str(args.output), 'slides': len(prs.slides),
                      'screenshot': str(args.screenshot),
                      'event_date': EVENT_DATE.date().isoformat(),
                      'template_sha256': TEMPLATE_SHA256}, ensure_ascii=False))


if __name__ == '__main__':
    main()
