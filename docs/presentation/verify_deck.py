#!/usr/bin/env python3
"""Verify editable deck/PDF parity and render every slide for visual review."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from pathlib import Path
from tempfile import NamedTemporaryFile

import fitz
from fontTools.ttLib import TTFont
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

ROOT = Path(__file__).resolve().parent
EVENT_DATE = '2026-09-29'
EVENT_ZIP_TIME = (2026, 9, 29, 12, 0, 0)
PDF_DATE = 'D:20260929120000Z'
PACKAGE_DATES = re.compile(r'\b\d{4}-\d{2}-\d{2}(?=T|\b)|\b\d{2}\.\d{2}\.\d{4}\b|D:\d{8}')
PROCESSING_LABELS = re.compile(r'фактическ(?:ая|и)\s+(?:генерация|сгенерировано|актуализация)'
                               r'|generated/verified|actual_(?:integration|generated|verified)', re.IGNORECASE)
TEAM_PLACEHOLDER = re.compile(
    r'\[[^\]]*(?:уточн\w*|фио|участник|@ник|телефон|город|регион|роль)[^\]]*\]'
    r'|\[\s*\]|\bуточнить\b', re.IGNORECASE,
)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def validate_date_text(text):
    """Package surfaces may carry only the organizer-authorized event date."""
    allowed = {EVENT_DATE, '29.09.2026', 'D:20260929'}
    dates = set(PACKAGE_DATES.findall(text))
    if dates - allowed or PROCESSING_LABELS.search(text):
        raise ValueError('Package contains a non-event date or processing timestamp label')


def validate_provenance(comments, pdf_subject):
    validate_date_text(comments)
    validate_date_text(pdf_subject)
    if '29.09.2026' not in comments or EVENT_DATE not in pdf_subject:
        raise ValueError('Organizer-authorized package date must be explicit')
    return {
        'package_date': EVENT_DATE,
        'event_only_metadata': 'PASS',
        'scope': 'Organizer-authorized package date; not proof of publication, integration or submission date',
    }


def stamp_pdf_event_date(path):
    """Full rewrite removes obsolete Info/XMP objects left by incremental saves."""
    path = Path(path)
    with NamedTemporaryFile(prefix='.event-date-', suffix='.pdf', dir=path.parent) as temporary:
        with fitz.open(path) as pdf:
            metadata = pdf.metadata
            metadata.update({
                'creationDate': PDF_DATE, 'modDate': PDF_DATE,
                'subject': 'Organizer-authorized package date ' + EVENT_DATE,
                'keywords': 'Thin Signals; LCT 2026; organizer-authorized event date',
            })
            pdf.set_metadata(metadata)
            pdf.del_xml_metadata()
            pdf.save(temporary.name, garbage=4, deflate=True, no_new_id=True)
        path.write_bytes(Path(temporary.name).read_bytes())


def validate_pptx_dates(path):
    with zipfile.ZipFile(path) as package:
        if package.comment:
            raise ValueError('Unexpected ZIP archive comment')
        for entry in package.infolist():
            if entry.date_time != EVENT_ZIP_TIME or entry.extra or entry.comment:
                raise ValueError('ZIP members must use the package date without extra timestamp fields')
            if (entry.filename.startswith(('docProps/', 'ppt/slides/', 'ppt/notesSlides/'))
                    and entry.filename.endswith('.xml')):
                validate_date_text(package.read(entry).decode('utf-8'))


def validate_pdf_dates(path):
    # Include raw obsolete objects and decompressed streams, not just current Info.
    validate_date_text(Path(path).read_bytes().decode('latin-1'))
    with fitz.open(path) as pdf:
        if pdf.metadata.get('creationDate') != PDF_DATE or pdf.metadata.get('modDate') != PDF_DATE:
            raise ValueError('PDF Info dates must equal the authorized package date')
        validate_date_text(pdf.get_xml_metadata())
        for xref in range(1, pdf.xref_length()):
            validate_date_text(pdf.xref_object(xref))
            if pdf.xref_is_stream(xref):
                validate_date_text(pdf.xref_stream(xref).decode('latin-1'))
        for page in pdf:
            validate_date_text(page.get_text())


def team_placeholders(slide_texts):
    """Report unresolved markers on team slides without requiring them to exist."""
    found = []
    for number, texts in slide_texts.items():
        if number not in (2, 3, 4):
            continue
        markers = [match.group(0) for text in texts for match in TEAM_PLACEHOLDER.finditer(text)]
        if markers:
            found.append({'slide': number, 'markers': markers})
    return found


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--deck', type=Path, default=ROOT / 'Thin-Signals-2026-09-29.pptx')
    parser.add_argument('--pdf', type=Path, default=ROOT / 'Thin-Signals-2026-09-29.pdf')
    parser.add_argument('--outdir', type=Path, required=True)
    parser.add_argument('--screenshot', type=Path, default=ROOT / 'assets/current-three-panels.png')
    parser.add_argument('--stamp-event-date', action='store_true')
    args = parser.parse_args()
    args.outdir.mkdir(parents=True, exist_ok=True)
    render = args.outdir / 'slides'
    render.mkdir(exist_ok=True)
    if args.stamp_event_date:
        stamp_pdf_event_date(args.pdf)
    validate_pptx_dates(args.deck)
    validate_pdf_dates(args.pdf)
    prs = Presentation(args.deck)
    pdf = fitz.open(args.pdf)
    issues = []
    assert len(prs.slides) == len(pdf) == 15
    assert prs.core_properties.created.date().isoformat() == '2026-09-29'
    assert prs.core_properties.modified.date().isoformat() == '2026-09-29'
    provenance = validate_provenance(
        prs.core_properties.comments, pdf.metadata.get('subject', ''),
    )
    with zipfile.ZipFile(args.deck) as zf:
        assert zf.testzip() is None
        assert len(zf.namelist()) == len(set(zf.namelist()))
        media = [x for x in zf.infolist() if x.filename.startswith('ppt/media/')]
        assert all(x.file_size > 0 for x in media)
        assert hashlib.sha256(args.screenshot.read_bytes()).digest() in [
            hashlib.sha256(zf.read(x.filename)).digest() for x in media
        ]
    signature = json.loads((ROOT / 'template-signature.json').read_text(encoding='utf-8'))
    for i in range(5):
        slide = prs.slides[i]
        assert slide.slide_layout.name == signature[i]['layout']
        assert [sh.shape_id for sh in slide.shapes] == signature[i]['shape_ids']
        for sh, (_, left, top, width, height) in zip(slide.shapes, signature[i]['geometry']):
            if i in (0, 1) and sh.shape_id == signature[i]['shape_ids'][1]:
                assert sh.is_placeholder
                continue
            assert (sh.left, sh.top, sh.width, sh.height) == (left, top, width, height)
    all_text = []
    slide_texts = {}
    for i, slide in enumerate(prs.slides):
        slide_texts[i + 1] = []
        for sh in slide.shapes:
            if sh.has_text_frame:
                all_text.append(sh.text)
                slide_texts[i + 1].append(sh.text)
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        assert r.font.name == 'Montserrat', (i + 1, r.text, r.font.name)
    pics = [s for s in prs.slides[13].shapes if hasattr(s, 'image')]
    assert len(pics) == 1
    pic = pics[0]
    assert (pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom) == (0, 0, 0, 0)
    with Image.open(args.screenshot) as im:
        assert abs(pic.width / pic.height - im.width / im.height) < 0.00001
    placeholders = team_placeholders(slide_texts)
    font_coverage = {}
    chars = {ord(c) for c in '\n'.join(all_text) if not c.isspace()}
    for style in ('Regular', 'Bold'):
        path = ROOT / 'fonts' / f'Montserrat-{style}.ttf'
        font = TTFont(path)
        missing = sorted(chars - set(font.getBestCmap()))
        assert not missing, (style, [chr(x) for x in missing])
        font_coverage[style] = {'sha256': digest(path), 'all_deck_characters_supported': True}
    pdf_fonts = {}
    for i, page in enumerate(pdf):
        assert abs(page.rect.width / page.rect.height - 16 / 9) < 0.001
        page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(render / f'slide-{i+1:02d}.png')
        normalized_pdf = re.sub(r'[^\w]', '', page.get_text()).casefold()
        for sh in prs.slides[i].shapes:
            if sh.has_text_frame and sh.text.strip():
                expected = re.sub(r'[^\w]', '', sh.text).casefold()
                if expected not in normalized_pdf:
                    issues.append({'slide': i+1, 'type': 'text_needs_visual_check', 'text': sh.text})
        for block in page.get_text('dict')['blocks']:
            for line in block.get('lines', []):
                for span in line['spans']:
                    x0, y0, x1, y1 = span['bbox']
                    if x0 < -1 or y0 < -1 or x1 > page.rect.width + 1 or y1 > page.rect.height + 1:
                        issues.append({'slide': i+1, 'type': 'text_outside_page', 'text': span['text']})
        for record in page.get_fonts(full=True):
            xref, _, _, family = record[:4]
            extracted = pdf.extract_font(xref)
            assert extracted[3], (i+1, family, 'font not embedded')
            pdf_fonts[family] = 'embedded'
    font = ImageFont.truetype(str(ROOT / 'fonts/Montserrat-Bold.ttf'), 20)
    sheet = Image.new('RGB', (1472, 1532), '#241431')
    draw = ImageDraw.Draw(sheet)
    for i in range(15):
        with Image.open(render / f'slide-{i+1:02d}.png') as slide:
            slide.thumbnail((468, 264))
            x, y = 16 + (i % 3) * 488, 16 + (i // 3) * 302
            sheet.paste(slide, (x, y))
            draw.text((x+4, y+268), f'{i+1:02d}', font=font, fill='white')
    sheet.save(args.outdir / 'presentation-contact-sheet.png')
    result = {
        'slides': len(prs.slides), 'pdf_pages': len(pdf), 'aspect': '16:9',
        'event_date': EVENT_DATE,
        'package_date_policy': provenance,
        'hidden_date_checks': 'PASS: core, notes, slides, ZIP members, PDF Info/XMP/raw objects/streams',
        'template_first_five': 'native layouts, order, shape IDs and non-logo geometry preserved',
        'all_editable_text_font': 'Montserrat',
        'editable_content': 'Native PowerPoint text and shapes; screenshot/template backgrounds are raster assets',
        'slide14_screenshot': {'sha256': digest(args.screenshot), 'crop': 'none', 'aspect_preserved': True},
        'font_coverage': font_coverage, 'pdf_fonts': pdf_fonts,
        'zip_integrity': 'PASS', 'nonempty_embedded_images': len(media),
        'automated_issues': issues,
        'pptx': {'bytes': args.deck.stat().st_size, 'sha256': digest(args.deck)},
        'pdf': {'bytes': args.pdf.stat().st_size, 'sha256': digest(args.pdf)},
        'team_fields': {
            'status': 'placeholders_detected' if placeholders else 'no_text_placeholders_detected',
            'placeholders': placeholders,
            'factual_validation': 'Not inferred from placeholder detection; supplied team facts need human verification',
        },
        'unresolved': (['team personal fields'] if placeholders else []) + ['independent end-to-end accuracy',
                       'authorized production UI smoke', 'complete durable evidence ledger',
                       'all-connectors-failed status contract', 'score monotonicity', 'competition submission'],
        'visual_review': 'required: inspect every slides/slide-NN.png; automated checks do not establish visual acceptance',
    }
    (args.outdir / 'presentation-verification.json').write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(1 if issues else 0)


if __name__ == '__main__':
    main()
